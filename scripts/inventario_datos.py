#!/usr/bin/env python3
"""Inventario de la informacion obtenida: dato, fuente y definicion de PYME.

Escribe pres/inventario_datos_pyme.xlsx con tres hojas:

  Inventario     una fila por dato: los tres campos pedidos (dato, fuente,
                 definicion de PYME usada) y, detras, si coincide con la
                 definicion estandar de la UE, el codigo de la columna
                 `criterio_segmentacion` de los CSV, el fichero, el periodo,
                 los paises y el numero de filas.
  Definiciones   que significa cada codigo (sale de schema.CRITERIOS) y los
                 umbrales de la Recomendacion 2003/361/CE.
  Fuentes        glosario de siglas: MIR, BSI, FM, SAFE, COREP, FINREP...

Periodo, paises y filas NO se escriben a mano: se leen de los CSV. Si una
fila del inventario apunta a un fichero que no existe, el script falla.

Uso:
    python3 scripts/inventario_datos.py
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import schema  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(ROOT, "pres", "inventario_datos_pyme.xlsx")

# Texto legible de cada criterio, para la columna "Definicion de PYME usada"
DEF = {
    "definicion_ue_2003_361": "Recomendación 2003/361/CE completa: menos de 250 empleados y hasta 50 M€ de facturación o 43 M€ de balance",
    "tamano_empresa_empleados": "Solo empleados: menos de 250",
    "pyme_crr_facturacion": "Solo facturación: hasta 50 M€ (art. 501 CRR)",
    "pyme_interna_banco": "Definición interna de cada banco (IRB): no armonizada",
    "pyme_mixta_estandar_e_irb": "Mezcla: estándar, facturación hasta 50 M€; IRB, definición interna de cada banco",
    "definicion_nacional": "Definición nacional de cada país: no armonizada",
    "tamano_prestamo": "Ninguna: segmenta por importe del PRÉSTAMO, no por tamaño de empresa",
    "sector_institucional": "Ninguna: persona física, sector hogares",
    "entidad": "No aplica: dato del banco entero",
    "n/a": "Ninguna: no segmenta por tamaño (todas las empresas, los hogares o el sistema)",
}
# Coincidencia con la definicion estandar de la UE
COINC = {
    "definicion_ue_2003_361": "Sí",
    "tamano_empresa_empleados": "Parcial: solo un criterio",
    "pyme_crr_facturacion": "Parcial: solo un criterio",
    "pyme_interna_banco": "No",
    "pyme_mixta_estandar_e_irb": "No",
    "definicion_nacional": "No",
    "tamano_prestamo": "No",
    "sector_institucional": "No aplica",
    "entidad": "No aplica",
    "n/a": "No aplica",
}

# (bloque, dato, fuente, codigo de criterio, fichero, matiz opcional)
INV = [
 # --- precio y producto ---
 ("Precio y producto", "Tipo de los préstamos a empresas por tramo de importe (≤0,25 M€, 0,25–1 M€, >1 M€), nueva producción, sin comisiones",
  "BCE, ECB Data Portal, dataset MIR", "tamano_prestamo", "prestamos_personales/prestamos_empresas_ecb_mir.csv",
  "Sociedades no financieras; deja fuera al autónomo"),
 ("Precio y producto", "Tipo medio anual ponderado por volumen, por tramo", "BCE, MIR (cálculo propio)",
  "tamano_prestamo", "prestamos_personales/prestamos_empresas_medias_ponderadas.csv", ""),
 ("Precio y producto", "Tipo de los préstamos a empresas, total de todas las cuantías", "BCE, ECB Data Portal, dataset MIR",
  "n/a", "prestamos_personales/prestamos_empresas_ecb_mir.csv", "Referencia de «todas las empresas» frente a la que se compara el circulante"),
 ("Precio y producto", "Prima PYME: tipo del préstamo pequeño menos el grande, dos definiciones", "BCE, MIR (cálculo propio)",
  "tamano_prestamo", "prestamos_personales/prima_pyme.csv", "≤0,25 M€ o ≤1 M€ frente a >1 M€"),
 ("Precio y producto", "Tipo y volumen del circulante (descubiertos y líneas de crédito)", "BCE, ECB Data Portal, dataset MIR",
  "n/a", "circulante/circulante_ecb_mir.csv", "El MIR no publica este producto por tramo de importe"),
 ("Precio y producto", "Cuña de comisiones (TAE menos tipo sin comisiones), tipo y volumen por tramo, España",
  "Banco de España, Boletín Estadístico, cap. 19", "tamano_prestamo", "prestamos_personales/bde_boletin.csv", ""),
 ("Precio y producto", "Tipo y saldo del circulante, España", "Banco de España, Boletín Estadístico, cap. 19",
  "n/a", "circulante/bde_circulante.csv", ""),
 ("Precio y producto", "Crédito a autónomos: tipo, TAE y volumen", "Banco de España, Boletín Estadístico, cap. 19",
  "sector_institucional", "prestamos_personales/bde_autonomos.csv", "El autónomo es persona física (hogares) en SEC 2010"),
 ("Precio y producto", "Tipo y volumen del crédito al consumo de hogares (proxy de autónomos)", "BCE, ECB Data Portal, dataset MIR",
  "n/a", "prestamos_personales/consumo_hogares_ecb_mir.csv", "Hogares, no empresas"),
 ("Precio y producto", "TAEG de préstamos a empresas, Italia", "Banca d'Italia, STACORIS, tavola TRI30951",
  "tamano_prestamo", "prestamos_personales/bancaditalia_taeg.csv", "Clase de importe disponible; excluye empresarios individuales"),
 ("Precio y producto", "Tipo de la línea de crédito o descubierto declarado por la empresa (Q8B)", "BCE, encuesta SAFE",
  "tamano_empresa_empleados", "circulante/safe_q8b.csv", "Serie hasta 2022-S1"),
 # --- coste de los recursos ---
 ("Coste de los recursos", "Coste del depósito de empresa: tipos a la vista y a plazo ponderados por saldos",
  "BCE, datasets MIR y BSI", "n/a", "liquidez/coste_recursos_pyme.csv", "Todas las sociedades no financieras (sector 2240)"),
 ("Coste de los recursos", "Coste del depósito de empresa, de hogares y del sistema", "BCE, datasets MIR y BSI (cálculo propio)",
  "n/a", "liquidez/coste_recursos_sistema.csv", ""),
 ("Coste de los recursos", "Tipos de depósito a plazo y a la vista, volumen", "BCE, ECB Data Portal, dataset MIR",
  "n/a", "liquidez/liquidez_ecb_mir.csv", "El MIR no desglosa los depósitos por importe"),
 ("Coste de los recursos", "Facilidad de depósito, operaciones principales y Euríbor a 3 meses", "BCE, ECB Data Portal, dataset FM",
  "n/a", "liquidez/tipos_oficiales_bce.csv", ""),
 # --- riesgo ---
 ("Riesgo", "PD, LGD, tasa de impago, tasa de pérdida y coste del riesgo PYME (PD × LGD)",
  "EBA, anexo de parámetros de riesgo (COREP C 9.02)", "pyme_interna_banco", "transversal/eba_parametros_riesgo.csv",
  "Solo cartera IRB; mediana de entidades declarantes"),
 ("Riesgo", "Mora y cobertura de la PYME (sociedades no financieras, of which SMEs)",
  "EBA, Risk Dashboard (FINREP)", "definicion_ue_2003_361", "transversal/eba_cartera.csv", "Muestra de bancos del EBA"),
 ("Riesgo", "Mora, cobertura y saldo del total, hogares, hipotecario y empresas",
  "EBA, Risk Dashboard (FINREP)", "n/a", "transversal/eba_cartera.csv", ""),
 ("Riesgo", "Saldo, mora y cobertura del inmueble comercial de empresas",
  "EBA, Risk Dashboard (FINREP)", "n/a", "hipotecas/eba_cartera.csv", ""),
 ("Riesgo", "Mora, NPE, refinanciaciones, coste del riesgo, eficiencia, ROE y CET1 del sistema",
  "EBA, Risk Dashboard", "n/a", "transversal/eba_indicadores.csv", "Grupo consolidado"),
 ("Riesgo", "Mora y cobertura de la cartera PYME en España, 5 bancos españoles",
  "EBA, EU-wide Transparency Exercise", "pyme_mixta_estandar_e_irb", "comparables_bancos/eba_te_bancos_es_serie.csv", ""),
 # --- capital ---
 ("Capital", "Densidad de RWA de la cartera PYME por país", "EBA, EU-wide Transparency Exercise",
  "pyme_mixta_estandar_e_irb", "transversal/eba_te_capital_pyme.csv", "Agregada por supervisor: incluye PYME extranjera"),
 ("Capital", "Densidad de RWA de PYME por método estándar", "EBA, EU-wide Transparency Exercise",
  "pyme_crr_facturacion", "transversal/eba_te_pyme_sa_irb.csv", ""),
 ("Capital", "Densidad de RWA de PYME por método IRB", "EBA, EU-wide Transparency Exercise",
  "pyme_interna_banco", "transversal/eba_te_pyme_sa_irb.csv", ""),
 ("Capital", "Cuota IRB, ahorro del IRB y descomposición de la brecha de densidad", "EBA, EU-wide Transparency Exercise (cálculo propio)",
  "pyme_mixta_estandar_e_irb", "transversal/eba_te_pyme_sa_irb.csv", ""),
 ("Capital", "Densidad de RWA de la PYME empresarial frente a la gran empresa", "EBA, EU-wide Transparency Exercise",
  "pyme_mixta_estandar_e_irb", "transversal/eba_te_capital_empresas.csv", ""),
 ("Capital", "Densidad descompuesta en PD, LGD y residuo con la fórmula IRB", "Cálculo propio: fórmula del art. 153 CRR con parámetros de COREP C 9.02",
  "pyme_interna_banco", "transversal/descomposicion_rwa.csv", "La fórmula usa además ventas del grupo < 50 M€ (supuesto: 25 M€)"),
 # --- demanda y acceso ---
 ("Demanda y acceso", "Tasa de rechazo, resultado de la solicitud y desánimo", "BCE, encuesta SAFE",
  "tamano_empresa_empleados", "transversal/safe_rechazo.csv", ""),
 ("Demanda y acceso", "Brecha de financiación (indicador de cambio, no de nivel)", "BCE, encuesta SAFE",
  "tamano_empresa_empleados", "transversal/safe_fg.csv", ""),
 ("Demanda y acceso", "Necesidad, destino, productos, obstáculos y bancarización, España",
  "CESGAR, XV Informe La financiación de la pyme en España", "tamano_empresa_empleados", "transversal/cesgar_demanda.csv",
  "Incluye personas físicas (autónomos)"),
 ("Demanda y acceso", "Rentabilidad del activo, coste de la deuda y su diferencia, por tamaño, España 1997–2024",
  "Banco de España, Central de Balances Integrada", "definicion_ue_2003_361", "transversal/cb_pymes.csv", ""),
 ("Demanda y acceso", "Spread PYME frente a gran empresa y tipo de interés", "OCDE, Financing SMEs and Entrepreneurs (Scoreboard)",
  "definicion_nacional", "prestamos_personales/oecd_scoreboard.csv", ""),
 ("Demanda y acceso", "Mora, retrasos de pago, rechazo y préstamos con aval público", "OCDE, Financing SMEs and Entrepreneurs (Scoreboard)",
  "definicion_nacional", "transversal/oecd_scoreboard.csv", ""),
 # --- factoring ---
 ("Factoring y confirming", "Volumen de factoring por país, cuota europea, penetración sobre PIB y variación", "EUF, EU Federation for Factoring",
  "n/a", "factoring_confirming/euf_factoring.csv", ""),
 ("Factoring y confirming", "Serie histórica de volumen de factoring desde 2007", "EUF, EU Turnover per country since 2007",
  "n/a", "factoring_confirming/euf_historico.csv", "Irlanda congelada desde 2021"),
 ("Factoring y confirming", "Factoring y confirming en España e Italia: cesiones y peso del confirming", "AEF y Assifact",
  "n/a", "factoring_confirming/factoring_nacional.csv", ""),
 ("Factoring y confirming", "Volumen de factoring", "OCDE, Financing SMEs and Entrepreneurs (Scoreboard)",
  "definicion_nacional", "factoring_confirming/oecd_scoreboard.csv", ""),
 # --- bancos ---
 ("Bancos", "Cuentas de segmento de 5 bancos europeos: ingresos, margen, comisiones, eficiencia, inversión",
  "Informes de resultados de Commerzbank, ABN AMRO, Intesa Sanpaolo, CaixaBank y BPER", "entidad",
  "comparables_bancos/comparables_bancos.csv", "Segmento de gestión: mezcla PYME y gran empresa"),
 ("Bancos", "Las mismas cuentas en % de la inversión crediticia", "Informes de resultados (cálculo propio)",
  "entidad", "comparables_bancos/comparables_ratios.csv", ""),
 ("Bancos", "Exposición, RWA, densidad y provisiones de la PYME española de 5 bancos españoles",
  "EBA, EU-wide Transparency Exercise", "pyme_mixta_estandar_e_irb", "comparables_bancos/eba_te_bancos_es.csv", "Contraparte en España"),
 ("Bancos", "CET1 y eficiencia de 5 bancos españoles", "EBA, EU-wide Transparency Exercise",
  "entidad", "comparables_bancos/eba_te_bancos_es.csv", ""),
 ("Bancos", "RoTE de grupo calculado y reconciliación con el ROE PYME", "EBA, EU-wide Transparency Exercise (cálculo propio)",
  "entidad", "comparables_bancos/reconciliacion_rote.csv", ""),
 # --- marco ---
 ("Marco institucional", "Información crediticia, derechos del acreedor, recuperación, tiempo y coste del concurso",
  "Banco Mundial, Doing Business 2020, más ordinales propios sobre normativa", "n/a", "transversal/marco_riesgo.csv", ""),
]

GLOSARIO = [
 ("MIR", "MFI Interest Rate Statistics (BCE)", "Tipos de interés de los bancos de la zona euro, nuevas operaciones y saldos, por producto, tramo de importe y plazo."),
 ("BSI", "Balance Sheet Items (BCE)", "Estadística de balance de los bancos de la zona euro: SALDOS de préstamos y depósitos por sector. Aquí se usa para ponderar los tipos de depósito del MIR por el peso real de la vista y del plazo."),
 ("FM", "Financial Markets (BCE)", "Tipos de mercado y oficiales: facilidad de depósito, operaciones principales, Euríbor."),
 ("SAFE", "Survey on the Access to Finance of Enterprises (BCE)", "Encuesta semestral a empresas sobre acceso a financiación."),
 ("COREP", "Common Reporting (EBA)", "Reporting supervisor de solvencia: exposiciones, RWA, PD y LGD. Reglamento de Ejecución (UE) 2021/451, anexo II."),
 ("FINREP", "Financial Reporting (EBA)", "Reporting supervisor contable: saldos, mora, cobertura. Reglamento de Ejecución (UE) 2021/451, anexo V."),
 ("Transparency Exercise", "EU-wide Transparency Exercise (EBA)", "Publicación anual del EBA con datos COREP y FINREP banco a banco."),
 ("Risk Dashboard", "Risk Dashboard (EBA)", "Publicación trimestral del EBA con indicadores de riesgo agregados por país."),
 ("CRR", "Capital Requirements Regulation", "Reglamento (UE) 575/2013 de requisitos de capital de los bancos."),
 ("IRB / estándar", "Métodos de cálculo del capital", "IRB: el banco usa sus propios modelos (PD, LGD) aprobados por el supervisor. Estándar: ponderaciones fijadas por la norma."),
 ("EUF", "EU Federation for Factoring", "Asociación europea del factoring; publica volúmenes por país."),
 ("CESGAR", "Confederación Española de Sociedades de Garantía", "Informe anual sobre la financiación de la pyme en España."),
]


def resumen_fichero(rel, cod):
    path = os.path.join(ROOT, rel)
    if not os.path.exists(path):
        raise SystemExit("el inventario apunta a un fichero que no existe: %s" % rel)
    with open(path, newline="", encoding="utf-8") as fh:
        filas = list(csv.DictReader(fh))
    # El codigo que el inventario atribuye al dato tiene que estar de verdad
    # en ese CSV; si no, el Excel diria una cosa y los datos otra.
    if cod not in {r["criterio_segmentacion"] for r in filas}:
        raise SystemExit("%s no contiene filas con criterio %r" % (rel, cod))
    pers = sorted(r["periodo_referencia"] for r in filas if r["periodo_referencia"])
    paises = sorted({r["pais"] for r in filas})
    return (len(filas), "%s a %s" % (pers[0], pers[-1]) if pers else "",
            len(paises))


def main():
    faltan = {c for _, _, _, c, _, _ in INV} - set(schema.CRITERIOS)
    if faltan:
        raise SystemExit("codigos que no estan en schema.CRITERIOS: %s" % faltan)

    wb = Workbook()
    F = "Arial"
    naranja = PatternFill("solid", fgColor="F56600")
    oscuro = PatternFill("solid", fgColor="2B2B2B")
    fino = Side(style="thin", color="CACACA")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    color_coinc = {"Sí": "C7E3F9", "Parcial: solo un criterio": "FFD6BA",
                   "No": "FFCCE8", "No aplica": "E5E5E5"}

    def cabecera(ws, cols, fills):
        for i, (c, fl) in enumerate(zip(cols, fills), 1):
            x = ws.cell(row=1, column=i, value=c)
            x.font = Font(name=F, bold=True, color="FFFFFF", size=10)
            x.fill = fl
            x.alignment = Alignment(vertical="center", wrap_text=True)
            x.border = borde
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 30

    # ---------------- hoja 1 ----------------
    ws = wb.active
    ws.title = "Inventario"
    cols = ["Dato", "Fuente", "Definición de PYME usada",
            "¿Coincide con la definición estándar de la UE?", "Matiz",
            "Bloque", "Código en los CSV (criterio_segmentacion)",
            "Fichero del repositorio", "Periodo", "Países", "Filas"]
    cabecera(ws, cols, [naranja] * 3 + [oscuro] * 8)
    for r, (bloque, dato, fuente, cod, fich, matiz) in enumerate(INV, 2):
        n, per, np_ = resumen_fichero(fich, cod)
        vals = [dato, fuente, DEF[cod], COINC[cod], matiz, bloque, cod,
                fich, per, np_, n]
        for c, v in enumerate(vals, 1):
            x = ws.cell(row=r, column=c, value=v)
            x.font = Font(name=F, size=9, bold=(c == 3))
            x.alignment = Alignment(vertical="top", wrap_text=True)
            x.border = borde
        ws.cell(row=r, column=4).fill = PatternFill(
            "solid", fgColor=color_coinc[COINC[cod]])
        ws.cell(row=r, column=11).number_format = "#,##0"
    for c, w in enumerate([46, 34, 42, 20, 34, 18, 26, 44, 20, 8, 9], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), len(INV) + 1)
    nota = len(INV) + 3
    ws.cell(row=nota, column=1, value=(
        "Las tres primeras columnas son las pedidas. Periodo, países y filas "
        "se leen de los CSV del repositorio al generar el fichero "
        "(scripts/inventario_datos.py); un mismo fichero puede aparecer en "
        "varias filas y entonces se repiten sus totales.")).font = Font(
            name=F, size=8, italic=True, color="818181")

    # ---------------- hoja 2 ----------------
    ws2 = wb.create_sheet("Definiciones")
    cabecera(ws2, ["Código", "Qué se llama PYME",
                   "¿Coincide con la estándar UE?", "Filas en los CSV"],
             [naranja] * 4)
    cuenta = {}
    import glob
    for f in glob.glob(os.path.join(ROOT, "*", "*.csv")):
        with open(f, newline="", encoding="utf-8") as fh:
            for x in csv.DictReader(fh):
                cuenta[x["criterio_segmentacion"]] = cuenta.get(
                    x["criterio_segmentacion"], 0) + 1
    for r, (cod, txt) in enumerate(schema.CRITERIOS.items(), 2):
        for c, v in enumerate([cod, txt, COINC[cod], cuenta.get(cod, 0)], 1):
            x = ws2.cell(row=r, column=c, value=v)
            x.font = Font(name=F, size=9, bold=(c == 1))
            x.alignment = Alignment(vertical="top", wrap_text=True)
            x.border = borde
        ws2.cell(row=r, column=3).fill = PatternFill(
            "solid", fgColor=color_coinc[COINC[cod]])
        ws2.cell(row=r, column=4).number_format = "#,##0"
    base = len(schema.CRITERIOS) + 4
    ws2.cell(row=base, column=1, value="Definición estándar de la UE (Recomendación 2003/361/CE)").font = Font(name=F, bold=True, size=10)
    tabla = [("", "Micro", "Pequeña", "Mediana"),
             ("Empleados", "< 10", "< 50", "< 250"),
             ("Facturación", "≤ 2 M€", "≤ 10 M€", "≤ 50 M€"),
             ("o balance", "≤ 2 M€", "≤ 10 M€", "≤ 43 M€")]
    for i, fila in enumerate(tabla):
        for j, v in enumerate(fila):
            x = ws2.cell(row=base + 1 + i, column=1 + j, value=v)
            x.font = Font(name=F, size=9, bold=(i == 0 or j == 0),
                          color="FFFFFF" if i == 0 else "2B2B2B")
            if i == 0:
                x.fill = oscuro
            x.border = borde
    ws2.cell(row=base + 6, column=1, value=(
        "Empleados siempre, más facturación O balance. En un grupo se suman "
        "las empresas asociadas y vinculadas. Fuente: Comisión Europea, "
        "DG GROW, página «SME definition».")).font = Font(
            name=F, size=8, italic=True, color="818181")
    for c, w in enumerate([30, 70, 26, 16], 1):
        ws2.column_dimensions[get_column_letter(c)].width = w

    # ---------------- hoja 3 ----------------
    ws3 = wb.create_sheet("Fuentes")
    cabecera(ws3, ["Sigla", "Nombre", "Qué es"], [naranja] * 3)
    for r, fila in enumerate(GLOSARIO, 2):
        for c, v in enumerate(fila, 1):
            x = ws3.cell(row=r, column=c, value=v)
            x.font = Font(name=F, size=9, bold=(c == 1))
            x.alignment = Alignment(vertical="top", wrap_text=True)
            x.border = borde
    for c, w in enumerate([20, 44, 90], 1):
        ws3.column_dimensions[get_column_letter(c)].width = w

    wb.save(SALIDA)
    print("%d datos -> %s" % (len(INV), SALIDA))
    return 0


if __name__ == "__main__":
    sys.exit(main())
