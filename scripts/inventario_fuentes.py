"""Inventario de fuentes publicas oficiales con informacion de PYME.

Para cada fuente: que publica sobre PYME, con que definicion, si la hemos
usado y para que. Despues, por componente del ROE: fuente usada, alternativas
y si era la mejor posible. Genera pres/inventario_fuentes_pyme.xlsx.

Verificaciones de disponibilidad hechas el 2026-09-28 contra las API
publicas (BCE Data Portal: CBD2, SUP, SAFE; Banque de France Webstat).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "pres", "inventario_fuentes_pyme.xlsx")

# ---------------------------------------------------------------------------
# 1. Fuentes
# ---------------------------------------------------------------------------
CAB_F = ["#", "Institución", "Fuente / dataset", "Qué publica sobre PYME",
         "Definición de PYME", "¿Coincide con la estándar UE?", "Cobertura",
         "Frecuencia y último dato", "¿La usamos?", "Para qué la usamos",
         "Qué tiene que NO usamos", "Acceso"]
F = [
 ["BCE", "MIR (MFI Interest Rate Statistics)",
  "Tipo de los préstamos nuevos a sociedades no financieras por tramo de importe (≤0,25 M€, 0,25–1 M€, >1 M€), sin comisiones. Tipos de depósito por sector.",
  "Ninguna: importe del PRÉSTAMO", "No (proxy)", "Todos los bancos de cada país",
  "Mensual · 2026-07", "Sí", "Precio del préstamo; coste de los depósitos",
  "—", "API abierta"],
 ["BCE", "BSI (Balance Sheet Items)",
  "Saldos de préstamos y depósitos por sector (empresas, hogares). No distingue tamaño.",
  "Ninguna: sector", "No aplica", "Todos los bancos", "Mensual · 2026-07", "Sí",
  "Ponderar los tipos de depósito (vista frente a plazo)", "—", "API abierta"],
 ["BCE", "SAFE (Survey on the Access to Finance of Enterprises)",
  "Encuesta a empresas: necesidad, solicitud, rechazo, brecha de financiación, obstáculos. Pregunta Q8B: tipo de interés que la empresa dice pagar por su línea de crédito o descubierto, PYME frente a gran empresa.",
  "Empleados <250", "Parcial (un criterio)", "Muestra de empresas, 7 países",
  "Semestral · 2025-S2 (Q8B solo hasta 2022-S1)", "Sí, salvo Q8B",
  "Rechazo, resultado de la solicitud, brecha",
  "Q8B, el único PRECIO por tamaño de EMPRESA armonizado. Se dejó de preguntar en 2022-S1. Valores 2022-S1 en la hoja «Evidencia»",
  "API abierta"],
 ["BCE", "BLS (Bank Lending Survey)",
  "Encuesta a bancos: criterios de concesión, márgenes y demanda del crédito a PYME frente a gran empresa.",
  "Facturación ≤50 M€ (definición del cuestionario)", "Parcial (un criterio)",
  "Muestra de bancos", "Trimestral", "No",
  "—", "Solo porcentajes netos (endurece / relaja): no da niveles de precio ni margen", "API abierta"],
 ["BCE", "Supervisory Banking Statistics (SUP)",
  "Entidades significativas por país: préstamos a PYME, ratio de mora PYME, préstamos PYME en fase 2. PD, LGD y ponderación de la cartera IRB de PYME, pero solo para el agregado del MUS.",
  "FINREP: Recomendación 2003/361/CE (mora); COREP IRB: interna (PD y LGD)", "Sí (mora); no (PD y LGD)",
  "Solo bancos significativos supervisados por el BCE", "Trimestral · 2026-Q2", "No",
  "—", "Mora y fase 2 de PYME un trimestre más recientes que el Risk Dashboard. PD y LGD ponderadas por exposición, pero no por país", "API abierta"],
 ["BCE", "CBD2 (Consolidated Banking Data)",
  "Todo el sistema bancario de cada país: eficiencia, ROE, margen de intereses. De PYME, solo la exposición IRB a empresas PYME (anual).",
  "COREP IRB: interna", "No", "Todo el sistema bancario (no solo una muestra)",
  "Trimestral / anual", "No", "—",
  "Eficiencia y márgenes del sistema completo, alternativa al Risk Dashboard. Las partidas de PD, LGD y ponderación PYME existen en el catálogo pero no tienen datos por país", "API abierta"],
 ["BCE / bancos centrales", "AnaCredit",
  "Préstamo a préstamo, con el tamaño de la empresa prestataria: tipo, importe, garantías, impago.",
  "Recomendación 2003/361/CE", "Sí", "Préstamos a personas jurídicas cuyo deudor tiene ≥25.000 € de riesgo total",
  "Mensual", "No", "—", "Sería la fuente ideal para todo el ROE por tamaño de empresa", "NO público: solo bancos centrales y supervisores"],
 ["EBA", "Risk Dashboard (anexo de datos, FINREP)",
  "Saldo, mora y cobertura de los préstamos a PYME. Del banco entero: eficiencia, coste del riesgo, margen de intereses, ROE, CET1.",
  "FINREP: Recomendación 2003/361/CE", "Sí (PYME); no aplica (banco entero)", "Muestra de bancos EBA, base consolidada",
  "Trimestral · 2026-Q1", "Sí", "Mora PYME, tamaño del mercado, eficiencia, CET1",
  "—", "Descarga abierta"],
 ["EBA", "EU-wide Transparency Exercise",
  "Por banco y país de contraparte: exposición, RWA, exposición en default y provisiones de PYME, por método estándar e IRB.",
  "COREP estándar: facturación ≤50 M€ (art. 501 CRR); COREP IRB: interna", "Parcial / no",
  "Bancos de la muestra EBA", "Anual · junio 2025", "Sí",
  "Densidad de RWA de PYME; cartera PYME de los 5 bancos españoles", "—", "Descarga abierta"],
 ["EBA", "Parámetros de riesgo (anexo, COREP C 9.02)",
  "PD y LGD de la cartera IRB de PYME por país.",
  "COREP IRB: interna de cada banco", "No", "Bancos IRB de la muestra EBA",
  "Trimestral · 2026-Q1", "Sí", "Coste del riesgo (PD × LGD)",
  "—", "Descarga abierta"],
 ["EBA", "Pillar 3 Data Hub",
  "Por banco: PD, LGD, exposición y RWA de la cartera PYME.",
  "COREP: estándar e interna", "Parcial / no", "Por banco", "Semestral", "No", "—",
  "Datos por banco. No accesible desde este entorno (el proxy lo bloquea)", "Web abierta"],
 ["Banco de España", "Boletín Estadístico, cap. 19",
  "Tipo sin comisiones (TEDR) y con comisiones (TAE) por tramo de importe; crédito a autónomos.",
  "Ninguna: importe del PRÉSTAMO", "No (proxy)", "España", "Mensual · 2026", "Sí",
  "Comisiones (cuña TAE − TEDR); autónomos", "—", "Descarga abierta"],
 ["Banco de España", "Central de Balances",
  "Rentabilidad del activo y coste de la deuda de las empresas por tamaño.",
  "Recomendación 2003/361/CE", "Sí", "Empresas españolas", "Anual · 2024", "Sí",
  "Capacidad de pago de la PYME", "—", "Descarga abierta"],
 ["Banca d'Italia", "STACORIS, tavola TRI30951",
  "TAEG de los préstamos de inversión por clase de importe (fuente AnaCredit).",
  "Ninguna: importe del préstamo", "No (proxy)", "Italia", "Trimestral · 2026-Q1",
  "Solo como referencia", "Confirmar la dirección de las comisiones",
  "No comparable con España (perímetro distinto)", "Descarga abierta"],
 ["Banque de France", "Webstat, «Crédits aux entreprises par taille»",
  "Tipo de los créditos nuevos por TAMAÑO DE EMPRESA: microempresa, PYME, ETI y gran empresa. Saldo del crédito a PYME. Mora de PYME (ACPR).",
  "Criterio LME (decreto 2008-1354): el de la Recomendación 2003/361/CE", "Sí",
  "Francia", "Mensual · catálogo actualizado en 2026-06", "No", "—",
  "Precio por tamaño de EMPRESA (no de préstamo) para Francia: permitiría contrastar el proxy del MIR. La API abierta publica el catálogo pero no los valores; el archivo descargable solo llega a 2013",
  "Web abierta; API de valores no abierta"],
 ["Bundesbank, Banco de Portugal, DNB, Central Bank of Ireland", "Estadísticas nacionales de tipos",
  "Verificado: ninguno publica tipo con comisiones para empresas. No verificado: si publican tipo por tamaño de empresa.",
  "—", "—", "Nacional", "—", "No", "—", "Pendiente de verificar el tipo por tamaño de empresa", "Web / API"],
 ["OCDE", "Financing SMEs and Entrepreneurs Scoreboard",
  "Diferencial de tipo PYME frente a gran empresa, tasa de rechazo, mora PYME, préstamos con aval público.",
  "Definición nacional de cada país", "No armonizada", "6 países (sin Alemania para el diferencial)",
  "Anual · 2022", "Sí", "Contraste del diferencial de precio", "—", "Descarga abierta"],
 ["CESGAR", "Informe de la financiación de la pyme en España",
  "Necesidad, destino, productos, obstáculos y resultado de la solicitud.",
  "Empleados <250, incluye autónomos", "Parcial (un criterio)", "España", "Anual · 2025", "Sí",
  "Demanda y obstáculos en España", "—", "Descarga abierta"],
 ["EUF, AEF, Assifact", "Estadísticas de factoring y confirming",
  "Volumen cedido por país; confirming en España e Italia.",
  "Ninguna: todas las empresas", "No aplica", "7 países", "Anual · 2025", "Sí",
  "Mercado de factoring", "—", "Descarga abierta"],
 ["Banco Mundial", "Doing Business 2020",
  "Información crediticia, derechos del acreedor, recuperación en concurso.",
  "Ninguna: marco del país", "No aplica", "7 países", "Mayo 2019 (discontinuado)", "Sí",
  "Entorno de información y recobro", "—", "Descarga abierta"],
 ["Registros públicos y mercantiles", "Normativa de CIRBE, Centrale dei Rischi, CRC, CCR, FIBEN, Millionenkredite y registros mercantiles",
  "Umbral de declaración al registro de crédito; acceso a las cuentas anuales de la PYME.",
  "Ninguna: marco del país", "No aplica", "7 países", "Normativa vigente 2026", "Sí",
  "Índice de información (elaboración propia)", "—", "Normativa pública"],
 ["Eurostat", "Structural Business Statistics",
  "Número de empresas, empleo y facturación por tamaño.",
  "Empleados (clases 0-9, 10-49, 50-249, 250+)", "Parcial (un criterio)", "UE", "Anual", "No", "—",
  "Tamaño del universo de PYME (demanda potencial)", "API abierta"],
 ["Bancos", "Informes de resultados",
  "Cuenta de resultados de los segmentos de negocio.",
  "Segmento de gestión de cada banco", "No", "5 bancos", "Semestral · 2026-H1", "Sí",
  "Orden de magnitud de comisiones y eficiencia de segmento", "—", "Web de cada banco"],
]

# ---------------------------------------------------------------------------
# 2. Componentes del ROE
# ---------------------------------------------------------------------------
CAB_C = ["Componente del ROE", "Fuente usada", "Definición de PYME de la fuente usada",
         "Alternativas públicas y su definición", "¿Era la mejor fuente posible?",
         "Mejora posible"]
C = [
 ["Precio del préstamo", "BCE, MIR: tramo ≤1 M€", "Importe del préstamo (proxy)",
  "SAFE Q8B: empresa <250 empleados, pero discontinuada en 2022-S1 y solo líneas de crédito, declarado por la empresa. Banque de France: tamaño de empresa, solo Francia. BdE y Banca d'Italia: importe, nacional. AnaCredit: tamaño de empresa, no público. EBA y BCE supervisor: no publican ingresos por intereses de PYME",
  "Sí. Es la única fuente armonizada y vigente para los 7 países. No existe un precio público actual por tamaño de empresa para los 7",
  "Contrastar el proxy con la serie francesa por tamaño de empresa (Banque de France) y con la SAFE Q8B hasta 2022"],
 ["Comisiones", "Banco de España: cuña TAE − TEDR (87 pb), aplicada a los 7", "Importe del préstamo",
  "Banca d'Italia TAEG (importe, perímetro no comparable). Los demás no publican tipo con comisiones para empresas",
  "Sí, pero solo es dato para España: en los otros 6 países es un supuesto", "Ninguna con dato público"],
 ["Coste de los recursos", "BCE: MIR (tipos de depósito de empresa) ponderado con BSI", "Sector (no aplica tamaño)",
  "—", "Sí. El coste de fondeo es del banco, no depende del tamaño del prestatario", "—"],
 ["Coste del riesgo (PD × LGD)", "EBA: parámetros de riesgo, COREP C 9.02, IRB PYME", "Interna de cada banco",
  "FINREP mora y cobertura de PYME (2003/361): mide el stock de morosos, no la pérdida esperada. Transparency Exercise: default y provisiones de PYME (estándar + IRB). BCE SUP: PD y LGD ponderadas por exposición, solo agregado del MUS",
  "Sí para una pérdida esperada por país. Límites: es la mediana de bancos (no ponderada por volumen) y solo cubre la cartera IRB",
  "Si el BCE publicara por país las PD y LGD ponderadas por exposición que ya tiene en su catálogo (SUP / CBD2), sustituirían a la mediana"],
 ["Capital: densidad de RWA", "EBA: Transparency Exercise", "Estándar: facturación ≤50 M€; IRB: interna",
  "Ninguna otra publica RWA de PYME por país (CBD2 solo la exposición IRB)",
  "Sí: es la única", "—"],
 ["Capital: ratio CET1", "EBA: Risk Dashboard", "Banco entero",
  "BCE CBD2 (sistema completo) y SUP (significativas)", "Sí; CBD2 cubriría todo el sistema en vez de la muestra EBA", "Opcional: CBD2"],
 ["Gastos de explotación", "EBA: Risk Dashboard, eficiencia", "Banco entero",
  "BCE CBD2 y SUP: también banco entero. Informes de bancos: solo ABN AMRO publica eficiencia de un segmento de empresas",
  "Sí; no existe la eficiencia de la banca PYME en fuente pública", "Opcional: CBD2 para cubrir todo el sistema"],
 ["Impuesto", "Tipo legal de sociedades de cada país", "No aplica", "—", "Sí", "—"],
 ["Volumen (ponderaciones, tamaño de mercado)", "EBA: Risk Dashboard, saldo PYME", "Recomendación 2003/361/CE",
  "BCE SUP (significativas, 2026-Q2); Transparency Exercise (exposición)", "Sí", "SUP da un trimestre más"],
]

# ---------------------------------------------------------------------------
# 3. Matriz fuente × componente
# ---------------------------------------------------------------------------
COMP = ["Precio", "Comisiones", "Coste de recursos", "PD y LGD", "Mora y cobertura",
        "RWA", "CET1", "Eficiencia", "Volumen"]
M = [
 ["BCE MIR", "Importe", "", "Sector", "", "", "", "", "", ""],
 ["BCE SAFE", "Empleados (hasta 2022)", "", "", "", "", "", "", "", ""],
 ["BCE SUP", "", "", "", "Interna (solo agregado MUS)", "2003/361", "", "Banco entero", "Banco entero", "2003/361"],
 ["BCE CBD2", "", "", "", "", "", "", "Banco entero", "Banco entero", ""],
 ["AnaCredit (no público)", "2003/361", "2003/361", "", "2003/361", "2003/361", "", "", "", "2003/361"],
 ["EBA Risk Dashboard", "", "", "", "", "2003/361", "", "Banco entero", "Banco entero", "2003/361"],
 ["EBA Transparency Exercise", "", "", "", "", "Mixta", "Mixta", "Banco entero", "", "Mixta"],
 ["EBA parámetros de riesgo", "", "", "", "Interna", "", "", "", "", ""],
 ["Banco de España", "Importe", "Importe", "", "", "", "", "", "", ""],
 ["Banca d'Italia", "Importe", "Importe", "", "", "", "", "", "", ""],
 ["Banque de France", "Tamaño empresa (solo FR)", "", "", "", "2003/361 (ACPR)", "", "", "", "2003/361"],
 ["OCDE", "Nacional (diferencial)", "", "", "", "Nacional", "", "", "", ""],
]
USADO = {("BCE MIR", "Precio"), ("BCE MIR", "Coste de recursos"),
         ("Banco de España", "Comisiones"), ("EBA parámetros de riesgo", "PD y LGD"),
         ("EBA Risk Dashboard", "Mora y cobertura"), ("EBA Transparency Exercise", "RWA"),
         ("EBA Risk Dashboard", "CET1"), ("EBA Risk Dashboard", "Eficiencia"),
         ("EBA Risk Dashboard", "Volumen")}

# ---------------------------------------------------------------------------
# 4. Evidencia descargada en la verificacion
# ---------------------------------------------------------------------------
PAISES = ["ES", "DE", "FR", "IT", "PT", "NL", "IE"]
Q8B = {"PYME": {"DE": 3.43, "ES": 2.45, "FR": 1.81, "IE": 6.04, "IT": 3.04, "NL": 2.91, "PT": 3.09},
       "Gran empresa": {"DE": 2.56, "ES": 1.71, "FR": 1.42, "IT": 1.32}}

# ---------------------------------------------------------------------------
NEGRO, NARANJA, GRIS = "2B2B2B", "F56600", "F2F4F6"
fina = Side(style="thin", color="D0D0D0")
borde = Border(left=fina, right=fina, top=fina, bottom=fina)


def hoja(ws, cab, filas, anchos):
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NEGRO)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = borde
    for f in filas:
        ws.append(f)
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = borde
    for i, a in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = a
    ws.freeze_panes = "B2"


wb = Workbook()
ws = wb.active
ws.title = "Léeme"
for linea in [
    "Inventario de fuentes públicas oficiales con información de PYME",
    "",
    "Fuentes: qué publica cada fuente sobre PYME, con qué definición y si la usamos.",
    "Componentes ROE: para cada pieza del ROE, qué fuente usamos, qué alternativas hay y si era la mejor posible.",
    "Matriz: fuente × componente; la celda dice con qué definición de PYME lo publica. En naranja, lo que usamos.",
    "Evidencia: datos descargados para verificar fuentes que no habíamos usado.",
    "",
    "Respuesta corta: ninguna fuente pública da todas las piezas del ROE de PYME. La familia más completa es",
    "la del supervisor (EBA: FINREP y COREP), que da volumen, mora, riesgo, capital, eficiencia y CET1, pero no",
    "precio, comisiones ni coste de recursos. Y aun dentro de la EBA hay tres definiciones de PYME distintas:",
    "FINREP (Recomendación 2003/361/CE), COREP estándar (facturación ≤50 M€) y COREP IRB (interna de cada banco).",
    "El precio por tamaño de empresa solo existe en AnaCredit (no público), en la SAFE hasta 2022 y en Francia.",
    "",
    "Verificaciones hechas el 2026-09-28 contra las API públicas del BCE (CBD2, SUP, SAFE) y del Banque de France (Webstat).",
]:
    ws.append([linea])
ws["A1"].font = Font(bold=True, size=14, color=NARANJA)
ws.column_dimensions["A"].width = 120

hoja(wb.create_sheet("Fuentes"), CAB_F, [[i + 1] + f for i, f in enumerate(F)],
     [4, 16, 28, 48, 30, 16, 22, 20, 12, 30, 44, 18])
hoja(wb.create_sheet("Componentes ROE"), CAB_C, C, [24, 30, 26, 60, 44, 40])

wm = wb.create_sheet("Matriz")
hoja(wm, ["Fuente"] + COMP, M, [26] + [18] * len(COMP))
for fila in wm.iter_rows(min_row=2):
    for j, c in enumerate(fila[1:]):
        if (fila[0].value, COMP[j]) in USADO:
            c.fill = PatternFill("solid", fgColor=NARANJA)
            c.font = Font(bold=True, color="FFFFFF")
        elif c.value:
            c.fill = PatternFill("solid", fgColor=GRIS)

we = wb.create_sheet("Evidencia")
we.append(["SAFE Q8B: tipo de interés de la línea de crédito o descubierto que declara la empresa, media ponderada, 2022-S1 (última encuesta con esta pregunta)"])
we["A1"].font = Font(bold=True)
we.append(["Tamaño"] + PAISES + ["Fuente"])
for t, v in Q8B.items():
    we.append([t] + [v.get(p, "sin dato") for p in PAISES] + ["BCE, SAFE, pregunta Q8B, SAFE_DENOM = WA"])
we.append([])
we.append(["Comprobación de disponibilidad de partidas PYME en el BCE (2026-09-28)"])
we["A5"].font = Font(bold=True)
for f in [
    ["SUP, I7008 (ratio de mora PYME)", "Con datos por país hasta 2026-Q2"],
    ["SUP, I7530 (préstamos PYME en fase 2)", "Con datos por país hasta 2026-Q2"],
    ["SUP, E0038 (préstamos PYME)", "Con datos por país hasta 2026-Q2"],
    ["SUP, EPD03 / EL003 / ERW03 (PD, LGD y ponderación IRB PYME)", "Solo agregado del MUS (B01), no por país"],
    ["CBD2, E3X31 (exposición IRB a empresas PYME)", "Anual, por país"],
    ["CBD2 y SUP, KSM08 / KSM10 / KS6_1 / KS6_2 (PD, LGD y ponderación PYME)", "En el catálogo, sin datos publicados"],
    ["Banque de France, pme-m-fr-ce-me-15-n-zz-pm (tipo crédito nuevo a PYME)", "Existe; la API abierta no devuelve valores"],
]:
    we.append(f)
we.column_dimensions["A"].width = 70
for col in "BCDEFGH":
    we.column_dimensions[col].width = 12
we.column_dimensions["I"].width = 40

wb.save(SALIDA)
print("escrito:", SALIDA)
